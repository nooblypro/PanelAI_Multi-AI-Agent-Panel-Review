"""File text extraction service.

Normalizes .pdf, .docx, and .txt files into clean plain text.
"""

from __future__ import annotations

import io
import logging
import os
import zipfile
import xml.etree.ElementTree as ET
from typing import Optional

from fastapi import UploadFile

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class FileExtractionError(Exception):
    """Raised when file content cannot be extracted or format is invalid."""


class UnsupportedFileTypeError(FileExtractionError):
    """Raised when an unsupported file extension is provided."""


def extract_text_from_bytes(content: bytes, filename: str) -> str:
    """Extract text from file bytes based on file extension.

    Parameters
    ----------
    content : bytes
        Raw file bytes.
    filename : str
        Name of the file including extension.

    Returns
    -------
    str
        Extracted plain text content.

    Raises
    ------
    UnsupportedFileTypeError
        If file extension is not .pdf, .docx, or .txt.
    FileExtractionError
        If file extraction fails due to corrupted content or parsing errors.
    """
    if not content:
        return ""

    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        return _extract_txt(content, filename)
    elif ext == ".pdf":
        return _extract_pdf(content, filename)
    elif ext == ".docx":
        return _extract_docx(content, filename)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}' for file '{filename}'. "
            f"Supported extensions: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )


async def extract_text_from_upload(upload_file: Optional[UploadFile]) -> str:
    """Extract plain text from a FastAPI UploadFile.

    Parameters
    ----------
    upload_file : Optional[UploadFile]
        FastAPI UploadFile object.

    Returns
    -------
    str
        Extracted plain text. Empty string if upload_file is None.
    """
    if upload_file is None or not upload_file.filename:
        return ""

    content = await upload_file.read()
    return extract_text_from_bytes(content, upload_file.filename)


# ---------------------------------------------------------------------------
# Internal format extractors
# ---------------------------------------------------------------------------


def _extract_txt(content: bytes, filename: str) -> str:
    """Extract text from .txt with multi-encoding fallback."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
    for enc in encodings:
        try:
            return content.decode(enc).strip()
        except (UnicodeDecodeError, LookupError):
            continue

    raise FileExtractionError(f"Could not decode text file '{filename}' with supported encodings.")


def _extract_pdf(content: bytes, filename: str) -> str:
    """Extract text from .pdf using pypdf."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages_text: list[str] = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                pages_text.append(page_text.strip())

        extracted = "\n\n".join(pages_text).strip()
        if not extracted and len(reader.pages) > 0:
            logger.warning("PDF '%s' has %d pages but extracted no text", filename, len(reader.pages))
        return extracted
    except Exception as exc:
        raise FileExtractionError(f"Failed to extract text from PDF '{filename}': {exc}") from exc


def _extract_docx(content: bytes, filename: str) -> str:
    """Extract text from .docx using python-docx with zipfile fallback."""
    # Attempt 1: python-docx
    try:
        import docx

        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs).strip()
    except Exception as docx_exc:
        logger.warning("python-docx extraction failed for '%s' (%s), trying XML fallback", filename, docx_exc)

    # Attempt 2: Direct word/document.xml extraction
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            xml_content = zf.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            # Find all text elements in WordprocessingML namespace
            text_nodes = tree.iter()
            extracted_words: list[str] = []
            for node in text_nodes:
                if node.tag.endswith("}t") and node.text:
                    extracted_words.append(node.text)
                elif node.tag.endswith("}p"):
                    extracted_words.append("\n")

            return "".join(extracted_words).strip()
    except Exception as xml_exc:
        raise FileExtractionError(
            f"Failed to extract text from DOCX '{filename}': {xml_exc}"
        ) from xml_exc
